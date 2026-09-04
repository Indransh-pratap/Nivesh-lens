import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.services.external.mf.portfolio_provider import SchemeHoldingItem


@dataclass(frozen=True)
class ParsedSchemePortfolio:
    scheme_code: str
    scheme_name: str
    ticker: str | None
    as_of_date: date
    holdings: list[SchemeHoldingItem]


def _parse_decimal(value: str) -> Decimal:
    cleaned = str(value).strip().replace(",", "").replace("%", "")
    if not cleaned:
        raise ValueError("Empty numeric value")

    try:
        return Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid numeric value: {value!r}") from exc


def _clean(value: str | None) -> str:
    if value is None:
        return ""

    return re.sub(r"\s+", " ", str(value)).strip()


def _find_column(headers: list[str], *names: str) -> int | None:
    normalized = [_clean(header).lower() for header in headers]

    for name in names:
        name = name.lower()

        for index, header in enumerate(normalized):
            if header == name or name in header:
                return index

    return None


def parse_portfolio_csv(content: str) -> list[SchemeHoldingItem]:
    """
    Parse a normalized AMC/AMFI portfolio disclosure CSV.
    """

    if not content.strip():
        return []

    reader = csv.reader(io.StringIO(content))
    rows = list(reader)

    if not rows:
        return []

    headers = [_clean(value) for value in rows[0]]

    company_index = _find_column(
        headers,
        "company name",
        "company",
        "name",
    )

    isin_index = _find_column(headers, "isin")

    sector_index = _find_column(
        headers,
        "sector",
        "industry",
    )

    weight_index = _find_column(
        headers,
        "weight (%)",
        "weight %",
        "weight",
        "% of net assets",
        "percentage",
    )

    if company_index is None:
        raise ValueError(
            "Portfolio disclosure is missing company-name column"
        )

    if weight_index is None:
        raise ValueError(
            "Portfolio disclosure is missing weight column"
        )

    holdings: list[SchemeHoldingItem] = []

    for row in rows[1:]:
        if not row:
            continue

        required_index = max(
            company_index,
            weight_index,
            isin_index or 0,
            sector_index or 0,
        )

        if len(row) <= required_index:
            continue

        company_name = _clean(row[company_index])

        if not company_name:
            continue

        try:
            weight = _parse_decimal(row[weight_index])
        except ValueError:
            continue

        if weight < 0:
            continue

        company_isin = (
            _clean(row[isin_index]).upper()
            if isin_index is not None
            else None
        )

        sector = (
            _clean(row[sector_index])
            if sector_index is not None
            else None
        )

        holdings.append(
            SchemeHoldingItem(
                company_name=company_name,
                company_isin=company_isin or None,
                sector=sector or None,
                weight_percentage=weight,
            )
        )

    return holdings


def _find_value(
    rows: list[tuple],
    label: str,
    value_offset: int = 1,
) -> object | None:
    """
    Find a metadata label and return the value immediately after it.
    """

    target = label.lower()

    for row in rows[:12]:
        for index, value in enumerate(row):
            if _clean(value).lower() == target:
                next_index = index + value_offset

                if next_index < len(row):
                    return row[next_index]

    return None


def _parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, str):
        value = value.strip()

        for fmt in (
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%Y-%m-%d",
        ):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                pass

    raise ValueError(f"Unable to parse portfolio date: {value!r}")


def _is_valid_isin(value: object) -> bool:
    isin = _clean(value).upper()

    return bool(
        re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", isin)
    )


def parse_portfolio_xlsx(
    content: bytes,
) -> list[ParsedSchemePortfolio]:
    """
    Parse an AMC/AMFI monthly portfolio XLSX workbook.

    Each worksheet normally represents one mutual-fund scheme.
    """

    if not content:
        return []

    workbook = load_workbook(
        filename=io.BytesIO(content),
        read_only=True,
        data_only=True,
    )

    portfolios: list[ParsedSchemePortfolio] = []

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))

        if not rows:
            continue

        # ---------------------------------------------------------
        # Scheme metadata
        # ---------------------------------------------------------

        scheme_code = _find_value(rows, "SBI Mutual Fund")

        # The AMC name is not always literally SBI, so find the
        # "SCHEME NAME :" label directly.
        scheme_name = _find_value(rows, "SCHEME NAME :")

        ticker = _find_value(rows, "SYMBOL / TICKER")

        portfolio_date = _find_value(
            rows,
            "PORTFOLIO STATEMENT AS ON :",
        )

        if scheme_code is None:
            # Generic AMC workbook:
            # look for a numeric scheme code near the top.
            for row in rows[:8]:
                for value in row:
                    text = _clean(value)

                    if re.fullmatch(r"\d{3,10}", text):
                        scheme_code = text
                        break

                if scheme_code is not None:
                    break

        if scheme_name is None or portfolio_date is None:
            # Not a portfolio-disclosure worksheet.
            continue

        scheme_code = _clean(scheme_code)

        if not scheme_code:
            continue

        as_of_date = _parse_date(portfolio_date)

        # ---------------------------------------------------------
        # Locate portfolio table header
        # ---------------------------------------------------------

        header_row_index: int | None = None

        for index, row in enumerate(rows):
            headers = [_clean(value).lower() for value in row]

            if (
                "name of the instrument / issuer" in headers
                and "isin" in headers
                and any("% to aum" in h for h in headers)
            ):
                header_row_index = index
                break

        if header_row_index is None:
            continue

        headers = [
            _clean(value)
            for value in rows[header_row_index]
        ]

        company_index = _find_column(
            headers,
            "name of the instrument / issuer",
        )

        isin_index = _find_column(headers, "isin")

        sector_index = _find_column(
            headers,
            "rating / industry",
            "industry",
        )

        weight_index = _find_column(
            headers,
            "% to aum",
            "% of aum",
        )

        if (
            company_index is None
            or isin_index is None
            or weight_index is None
        ):
            continue

        # ---------------------------------------------------------
        # Parse equity holdings only.
        # ---------------------------------------------------------

        holdings: list[SchemeHoldingItem] = []

        in_equity_section = False

        for row in rows[header_row_index + 1:]:
            values = list(row)

            company = (
                _clean(values[company_index])
                if company_index < len(values)
                else ""
            )

            isin = (
                _clean(values[isin_index]).upper()
                if isin_index < len(values)
                else ""
            )

            sector = (
                _clean(values[sector_index])
                if sector_index is not None
                and sector_index < len(values)
                else ""
            )

            weight_value = (
                values[weight_index]
                if weight_index < len(values)
                else None
            )

            company_upper = company.upper()

            # Enter equity section.
            if "EQUITY & EQUITY RELATED" in company_upper:
                in_equity_section = True
                continue

            # Stop stock look-through once debt begins.
            if "DEBT INSTRUMENTS" in company_upper:
                in_equity_section = False
                continue

            if not in_equity_section:
                continue

            # Ignore headings/categories.
            if not company or not _is_valid_isin(isin):
                continue

            try:
                weight = _parse_decimal(weight_value)
            except (ValueError, TypeError):
                continue

            if weight < 0:
                continue

            holdings.append(
                SchemeHoldingItem(
                    company_name=company,
                    company_isin=isin,
                    sector=sector or None,
                    weight_percentage=weight,
                )
            )

        portfolios.append(
            ParsedSchemePortfolio(
                scheme_code=scheme_code,
                scheme_name=_clean(scheme_name),
                ticker=_clean(ticker) or None,
                as_of_date=as_of_date,
                holdings=holdings,
            )
        )

    workbook.close()

    return portfolios